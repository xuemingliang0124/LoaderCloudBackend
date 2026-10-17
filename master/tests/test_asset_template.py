"""资产导入模板下载契约测试：格式矩阵 + 模板表头与列映射一致性 + 上传回环。

核心价值：回环用例（下载模板 → 上传 → 解析抽取成功）端到端证明
「模板表头 ↔ asset_parser 列映射规则」严格一致（模板动态生成的初衷）。
"""

import json as jsonlib
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.security import create_access_token
from app.services.asset_parser import (
    ENV_INVENTORY_COLUMN_MAP,
    TXN_INVENTORY_COLUMN_MAP,
    match_column,
    parse_markdown,
    run_asset_parse,
)


class _FakeStore:
    def __init__(self) -> None:
        self.upserted: list = []

    async def upsert_chunks(self, points) -> None:
        self.upserted.extend(points)


class _FakeEmbeddings:
    async def aembed_documents(self, texts):
        return [[0.1, 0.2, 0.3, 0.4] for _ in texts]

    async def aembed_query(self, text):
        return [0.1, 0.2, 0.3, 0.4]


@pytest.fixture
def fake_vector_store(monkeypatch) -> _FakeStore:
    store = _FakeStore()
    monkeypatch.setattr(
        "app.services.langchain_pipeline.get_vector_store", lambda: store
    )
    return store


@pytest.fixture
def fake_embeddings(monkeypatch) -> _FakeEmbeddings:
    embeddings = _FakeEmbeddings()
    monkeypatch.setattr(
        "app.services.langchain_pipeline.get_embeddings", lambda: embeddings
    )
    return embeddings


def _auth(username: str, role: str = "user") -> dict:
    return {"Authorization": f"Bearer {create_access_token(username, role)}"}


async def _create_project(client, name: str, username: str = "alice") -> int:
    resp = await client.post(
        "/api/v1/projects", json={"name": name}, headers=_auth(username)
    )
    assert resp.status_code == 200
    return resp.json()["data"]["id"]


def _patch_storage(monkeypatch) -> None:
    _FAKE_FILES: dict[str, bytes] = {}

    async def _fake_get(object_key: str):
        return _FAKE_FILES[object_key]

    async def _fake_upload(object_key, data, content_type=""):
        _FAKE_FILES[object_key] = data

    monkeypatch.setattr("app.services.storage.upload_bytes", _fake_upload)
    monkeypatch.setattr("app.services.asset_parser.storage.upload_bytes", _fake_upload)
    monkeypatch.setattr("app.services.asset_parser.storage.get_object_bytes", _fake_get)


def _assert_headers_match_column_map(headers: list[str], asset_type: str) -> None:
    """模板表头逐列核对：必须命中对应列映射，且必填列（编码）在列。"""
    column_map = (
        ENV_INVENTORY_COLUMN_MAP
        if asset_type == "env_inventory"
        else TXN_INVENTORY_COLUMN_MAP
    )
    for header in headers:
        assert match_column(header, column_map), f"表头 {header!r} 未命中列映射"
    required = "环境编码" if asset_type == "env_inventory" else "交易码"
    assert required in headers


# ---------- 门禁与错误码 ----------


async def test_template_requires_token() -> None:
    from httpx import ASGITransport, AsyncClient

    from app.main import app

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        resp = await client.get("/api/v1/assets/templates/env_inventory")
    assert resp.status_code == 401


async def test_template_unknown_asset_type_3060(client) -> None:
    resp = await client.get(
        "/api/v1/assets/templates/not_a_type", headers=_auth("alice")
    )
    assert resp.status_code == 400
    assert resp.json()["code"] == 3060


async def test_template_format_not_supported_3060(client) -> None:
    """文档类模板请求 xlsx → 3060（文档类仅 md）。"""
    resp = await client.get(
        "/api/v1/assets/templates/plan_doc",
        params={"fmt": "xlsx"},
        headers=_auth("alice"),
    )
    assert resp.status_code == 400
    assert resp.json()["code"] == 3060
    assert "md" in resp.json()["message"]


# ---------- 三种格式内容校验 ----------


async def test_env_inventory_xlsx_template(client) -> None:
    resp = await client.get(
        "/api/v1/assets/templates/env_inventory",
        params={"fmt": "xlsx"},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]

    import io

    workbook = load_workbook(io.BytesIO(resp.content))
    sheet = workbook["清单"]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    _assert_headers_match_column_map(headers, "env_inventory")
    # 第二行为示例行（非空）
    assert any(c.value for c in next(sheet.iter_rows(min_row=2, max_row=2)))
    # 填写说明 sheet 存在
    assert "填写说明" in workbook.sheetnames


async def test_txn_inventory_json_template(client) -> None:
    resp = await client.get(
        "/api/v1/assets/templates/txn_inventory",
        params={"fmt": "json"},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/json")

    payload = jsonlib.loads(resp.content)
    assert set(payload.keys()) == {"items"}  # 严格 {"items": [...]} 形态
    assert isinstance(payload["items"], list) and payload["items"]
    headers = list(payload["items"][0].keys())
    _assert_headers_match_column_map(headers, "txn_inventory")


async def test_txn_inventory_md_template(client) -> None:
    resp = await client.get(
        "/api/v1/assets/templates/txn_inventory",
        params={"fmt": "md"},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/markdown")

    parsed = parse_markdown(resp.content)
    assert parsed.tables, "md 清单模板应包含管道表格"
    _assert_headers_match_column_map(parsed.tables[0][0], "txn_inventory")


async def test_plan_doc_md_template(client) -> None:
    resp = await client.get(
        "/api/v1/assets/templates/plan_doc",
        params={"fmt": "md"},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200
    text = resp.content.decode("utf-8")
    assert "# 性能测试方案" in text


# ---------- 回环：下载模板 → 上传 → 解析抽取成功 ----------


async def test_env_json_template_upload_roundtrip(client, db_env, monkeypatch) -> None:
    """下载 json 模板原样上传 → 解析 → 环境抽取成功（模板与解析规则一致的端到端证明）。"""
    _patch_storage(monkeypatch)
    template = (
        await client.get(
            "/api/v1/assets/templates/env_inventory",
            params={"fmt": "json"},
            headers=_auth("alice"),
        )
    ).content

    pid = await _create_project(client, "模板回环项目")
    resp = await client.post(
        f"/api/v1/projects/{pid}/assets",
        data={"asset_type": "env_inventory"},
        files={"file": ("环境清单.json", template, "application/json")},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200, resp.text
    asset_id = resp.json()["data"]["id"]
    await run_asset_parse(asset_id, db_env)

    from sqlalchemy import select

    from app.models.asset import Asset
    from app.models.environment import Environment

    async with db_env() as db:
        row = (await db.execute(select(Asset).where(Asset.id == asset_id))).scalar_one()
        assert row.status == "ready", row.parse_meta
        assert row.parse_meta["environments"] == 1
        env = (
            await db.execute(select(Environment).where(Environment.project_id == pid))
        ).scalar_one()
    assert env.env_code == "PROD"


async def test_env_md_template_upload_roundtrip(
    client, db_env, monkeypatch, fake_vector_store, fake_embeddings
) -> None:
    """下载 md 模板原样上传 → 解析 → 结构化抽取 + 说明文字入向量。"""
    _patch_storage(monkeypatch)
    monkeypatch.setattr(
        "app.services.langchain_pipeline.get_settings",
        lambda: SimpleNamespace(
            embedding_provider="",
            embedding_base_url="",
            embedding_api_key="",
            embedding_model="",
            embedding_batch_size=64,
            embedding_timeout=5,
            embedding_max_retries=1,
            embedding_dims=4,
            chunk_size=500,
            chunk_overlap=50,
            timezone="Asia/Shanghai",
        ),
    )
    template = (
        await client.get(
            "/api/v1/assets/templates/env_inventory",
            params={"fmt": "md"},
            headers=_auth("alice"),
        )
    ).content

    pid = await _create_project(client, "md模板回环项目")
    resp = await client.post(
        f"/api/v1/projects/{pid}/assets",
        data={"asset_type": "env_inventory"},
        files={"file": ("环境清单.md", template, "text/markdown")},
        headers=_auth("alice"),
    )
    assert resp.status_code == 200, resp.text
    asset_id = resp.json()["data"]["id"]
    await run_asset_parse(asset_id, db_env)

    from sqlalchemy import select

    from app.models.asset import Asset
    from app.models.environment import Environment

    async with db_env() as db:
        row = (await db.execute(select(Asset).where(Asset.id == asset_id))).scalar_one()
        assert row.status == "ready", row.parse_meta
        assert row.parse_meta["environments"] == 1
        assert row.parse_meta["chunks"] >= 1  # 模板说明文字入知识库
        env = (
            await db.execute(select(Environment).where(Environment.project_id == pid))
        ).scalar_one()
    assert env.env_code == "PROD"
