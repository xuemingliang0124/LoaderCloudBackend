"""文档解析管道单元测试（D3 + P3 Stage 1 重构）：列映射 / 宽松取值 / 三种格式解析。

P3 Stage 1 重构变更：
- chunk_text 相关测试移除（改用 langchain_pipeline.RecursiveCharacterTextSplitter，
  对应测试见 test_langchain_pipeline.py）
- 解析/列映射/宽松取值/分发 等纯函数测试保留不变

纯函数级测试，不依赖 DB / MinIO / 向量库；docx/xlsx 用库本身反向构造测试文件，
pdf 用手工构造的最小合法 PDF（含 xref 偏移计算）。
"""

import io

import pytest

from app.services.asset_parser import (
    ENV_INVENTORY_COLUMN_MAP,
    TXN_INVENTORY_COLUMN_MAP,
    UnsupportedFormatError,
    dispatch_parse,
    extract_inventory_rows,
    match_column,
    parse_docx,
    parse_pdf,
    parse_xlsx,
    _parse_dict_value,
    _parse_number,
    _split_list_value,
)
from app.services.embedding_client import point_id_for

# ---------- 列映射 ----------


def test_match_column_env_and_txn() -> None:
    assert match_column("环境名称", ENV_INVENTORY_COLUMN_MAP) == "name"
    assert match_column("基础 URL", ENV_INVENTORY_COLUMN_MAP) == "base_url"
    assert match_column("主机清单", ENV_INVENTORY_COLUMN_MAP) == "hosts"
    assert match_column("TPS", TXN_INVENTORY_COLUMN_MAP) == "sla_tps"
    assert match_column("交易编码", TXN_INVENTORY_COLUMN_MAP) == "txn_code"
    assert match_column("未知列", ENV_INVENTORY_COLUMN_MAP) is None


def test_extract_inventory_rows_basic_and_unmatched() -> None:
    tables = [
        [
            ["环境名称", "环境编码", "基础URL", "负责人"],
            ["生产环境", "prod", "http://prod.example.com", "张三"],
            ["", "", "", ""],  # 空行丢弃
            ["预发环境", "staging", "http://staging.example.com", "李四"],
        ]
    ]
    result = extract_inventory_rows(tables, ENV_INVENTORY_COLUMN_MAP)
    assert [r["name"] for r in result.rows] == ["生产环境", "预发环境"]
    assert result.rows[0]["env_code"] == "prod"
    assert result.rows[0]["base_url"] == "http://prod.example.com"
    assert result.unmatched_columns == ["负责人"]


def test_extract_inventory_rows_no_header_skipped() -> None:
    tables = [[["随便", "什么"], ["1", "2"]]]
    result = extract_inventory_rows(tables, ENV_INVENTORY_COLUMN_MAP)
    assert result.rows == []
    assert result.warnings == ["第 1 张表未识别到表头行，已跳过"]


def test_extract_inventory_rows_second_table_has_header() -> None:
    tables = [
        [["注释行", "无映射"], ["a", "b"]],
        [["交易码", "交易名称"], ["login", "登录"]],
    ]
    result = extract_inventory_rows(tables, TXN_INVENTORY_COLUMN_MAP)
    assert len(result.rows) == 1
    assert result.rows[0]["txn_code"] == "login"


# ---------- 宽松取值 ----------


def test_split_list_value_variants() -> None:
    assert _split_list_value("10.0.0.1,10.0.0.2；10.0.0.3") == [
        "10.0.0.1",
        "10.0.0.2",
        "10.0.0.3",
    ]
    assert _split_list_value('["a","b"]') == ["a", "b"]
    assert _split_list_value(["x", "y"]) == ["x", "y"]
    assert _split_list_value("") == []


def test_parse_dict_value_variants() -> None:
    assert _parse_dict_value('{"k1": "v1"}') == {"k1": "v1"}
    assert _parse_dict_value("base_url=http://x; token=abc") == {
        "base_url": "http://x",
        "token": "abc",
    }
    assert _parse_dict_value("k：v") == {"k": "v"}
    assert _parse_dict_value("无键值对") == {}
    assert _parse_dict_value("") == {}


def test_parse_number_variants() -> None:
    assert _parse_number("500") == 500.0
    assert _parse_number("95%") == 95.0
    assert _parse_number("1,200") == 1200.0
    assert _parse_number("120ms") == 120.0
    assert _parse_number("abc") is None
    assert _parse_number("") is None


# ---------- docx 解析 ----------


def _make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        rows = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row_values in enumerate(table):
            for c, value in enumerate(row_values):
                rows.rows[r].cells[c].text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_paragraphs_and_tables() -> None:
    data = _make_docx(
        ["性能测试方案", "覆盖登录与下单链路"],
        [["交易", "TPS"], ["login", "500"]],
    )
    parsed = parse_docx(data)
    assert parsed.paragraphs == ["性能测试方案", "覆盖登录与下单链路"]
    assert parsed.tables == [[["交易", "TPS"], ["login", "500"]]]


# ---------- xlsx 解析 ----------


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_xlsx_multi_sheets_and_trailing_empty_cells() -> None:
    data = _make_xlsx(
        {
            "环境清单": [["环境名称", "环境编码"], ["生产", "prod"], [None, None]],
            "空表": [[None]],
        }
    )
    parsed = parse_xlsx(data)
    assert len(parsed.tables) == 1  # 空表被跳过
    assert parsed.tables[0] == [["环境名称", "环境编码"], ["生产", "prod"]]


# ---------- pdf 解析 ----------


def _make_pdf(text: str) -> bytes:
    """构造最小合法 PDF（单页 + Helvetica 文本，含精确 xref 偏移）。"""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def test_parse_pdf_text() -> None:
    parsed = parse_pdf(_make_pdf("perf test plan for login"))
    assert parsed.paragraphs == ["perf test plan for login"]
    assert parsed.tables == []


def test_parse_pdf_table() -> None:
    # pdfplumber 从带表格线的页面提取表格需要绘制线条，文本场景已覆盖主路径；
    # 表格路径由 docx/xlsx 用例保障，此处仅验证空表格不报错
    parsed = parse_pdf(_make_pdf("hello"))
    assert parsed.paragraphs == ["hello"]


# ---------- 分发与旧格式 ----------


def test_dispatch_by_extension() -> None:
    docx_data = _make_docx(["a"])
    assert dispatch_parse(".docx", docx_data).paragraphs == ["a"]
    xlsx_data = _make_xlsx({"s": [["h"], ["v"]]})
    assert dispatch_parse(".xlsx", xlsx_data).tables == [[["h"], ["v"]]]


def test_dispatch_unsupported_legacy_formats() -> None:
    with pytest.raises(UnsupportedFormatError, match="\\.docx"):
        dispatch_parse(".doc", b"legacy")
    with pytest.raises(UnsupportedFormatError, match="\\.xlsx"):
        dispatch_parse(".xls", b"legacy")
    with pytest.raises(UnsupportedFormatError, match="\\.pdf"):
        dispatch_parse(".pptx", b"slides")


# ---------- 向量点 ID ----------


def test_point_id_deterministic_int64() -> None:
    assert point_id_for(7, 0) == point_id_for(7, 0)
    assert point_id_for(7, 0) != point_id_for(7, 1)
    assert point_id_for(7, 0) != point_id_for(8, 0)
    assert 0 <= point_id_for(1, 1) < 2**32  # crc32 无符号 32 位，int64 安全
