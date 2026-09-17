"""文档解析管道（D3，roadmap 核心）：上传资产 → 解析 → 双路输出。

双路输出：
- (a) 文本切片（300-500 字/段，overlap 50 字）→ embedding（D4）→ Qdrant
  （经 VectorStore 协议 upsert_chunks，业务层禁止 import Qdrant SDK）
- (b) 表格行 → 列映射规则（下方常量）→ test_environment / test_transaction 表

关键设计（ptp-dev 3.1 强约束）：
- 全部 CPU 密集解析用 asyncio.to_thread 包裹，禁止事件循环内阻塞
- 按扩展名分发：docx→python-docx / xlsx→openpyxl / pdf→pdfplumber
  （.doc/.xls/.pptx 为旧格式，解析器不支持 → FAILED 并提示另存）
- 状态机：PENDING → PARSING → READY / FAILED；FAILED 由 retry-parse 端点重新入队
- 上传后由 APScheduler 延迟投递（scheduler.enqueue_date_job），不阻塞上传响应
- 状态迁移全部用 Core update() 语句（避免异步会话身份映射/懒加载陷阱，
  见 project_memory 异步会话陷阱三连）
- 未配置 Embedding 时降级：结构化抽取照常，仅跳过向量入库（parse_meta.indexed=false）

parse_meta 结构（AssetOut 透出，D5 列映射修正接口消费）：
{"chunks": 切片数, "indexed": 是否入向量库, "environments": 抽取环境行数,
 "transactions": 抽取交易行数, "warnings": [跳过原因], "unmatched_columns": [未匹配表头],
 "error": 失败原因（仅 FAILED）}
"""

import asyncio
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from loguru import logger
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.db.session import SessionLocal
from app.models.asset import Asset
from app.models.enums import AssetStatus, AssetType
from app.models.environment import Environment
from app.models.script import Script
from app.models.transaction import Transaction
from app.services import storage
from app.services.embedding_client import EmbeddingClient, point_id_for
from app.services.vector_store import ChunkDoc, get_vector_store

# 文本切片参数（roadmap D3：300-500 字/段，overlap 50 字）
CHUNK_TARGET_SIZE = 500
CHUNK_OVERLAP = 50

# 结构化抽取的资产类型 → 列映射表（D5 remap 前的硬编码兜底）
# 键为归一化表头（小写、去空格）；命中列才进入抽取，未命中的记入 unmatched_columns
ENV_INVENTORY_COLUMN_MAP: dict[str, str] = {
    "环境名称": "name",
    "环境名": "name",
    "名称": "name",
    "环境编码": "env_code",
    "环境标识": "env_code",
    "编码": "env_code",
    "基础url": "base_url",
    "基础地址": "base_url",
    "服务地址": "base_url",
    "url": "base_url",
    "地址": "base_url",
    "主机": "hosts",
    "主机清单": "hosts",
    "主机列表": "hosts",
    "数据库连接": "db_connections",
    "数据库": "db_connections",
    "中间件": "middleware_info",
    "中间件信息": "middleware_info",
    "变量": "variables",
    "变量覆盖": "variables",
    "参数": "variables",
    "参数覆盖": "variables",
    "描述": "description",
    "说明": "description",
    "备注": "description",
}
TXN_INVENTORY_COLUMN_MAP: dict[str, str] = {
    "交易名称": "name",
    "交易名": "name",
    "名称": "name",
    "交易码": "txn_code",
    "交易编码": "txn_code",
    "交易编号": "txn_code",
    "目标tps": "sla_tps",
    "sla_tps": "sla_tps",
    "tps": "sla_tps",
    "吞吐量": "sla_tps",
    "p95": "sla_p95_ms",
    "p95(ms)": "sla_p95_ms",
    "p95响应时间": "sla_p95_ms",
    "响应时间(ms)": "sla_p95_ms",
    "错误率": "sla_error_rate",
    "sla错误率": "sla_error_rate",
    "默认脚本": "default_script",
    "默认脚本名": "default_script",
    "描述": "description",
    "说明": "description",
    "备注": "description",
}

_INVENTORY_COLUMN_MAPS: dict[str, dict[str, str]] = {
    AssetType.ENV_INVENTORY.value: ENV_INVENTORY_COLUMN_MAP,
    AssetType.TXN_INVENTORY.value: TXN_INVENTORY_COLUMN_MAP,
}


class UnsupportedFormatError(RuntimeError):
    """解析器不支持的文件格式（旧格式 .doc/.xls/.pptx）。"""


@dataclass
class ParsedDoc:
    """解析中间产物：段落文本（含 pdf 按页文本）+ 表格（每表为二维文本行）。"""

    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)


@dataclass
class InventoryExtraction:
    """结构化抽取结果：行字典列表 + 未匹配表头 + 警告。"""

    rows: list[dict[str, str]] = field(default_factory=list)
    unmatched_columns: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# ---------- 格式解析（同步 CPU 密集，统一由 asyncio.to_thread 调用） ----------


def parse_docx(data: bytes) -> ParsedDoc:
    """python-docx：非空段落 + 表格（单元格文本）。"""
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(c for c in r)]
        if rows:
            tables.append(rows)
    return ParsedDoc(paragraphs=paragraphs, tables=tables)


def parse_xlsx(data: bytes) -> ParsedDoc:
    """openpyxl：每个非空工作表视为一张表（read_only + data_only 取计算值）。"""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables: list[list[list[str]]] = []
    for worksheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                rows.append(cells)
        if rows:
            tables.append(rows)
    return ParsedDoc(tables=tables)


def parse_pdf(data: bytes) -> ParsedDoc:
    """pdfplumber：按页提取文本 + 表格。"""
    import pdfplumber

    paragraphs: list[str] = []
    tables: list[list[list[str]]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = (page.extract_text() or "").strip()
            if text:
                paragraphs.append(text)
            for table in page.extract_tables() or []:
                rows = [[(cell or "").strip() for cell in row] for row in table if row]
                rows = [r for r in rows if any(r)]
                if rows:
                    tables.append(rows)
    return ParsedDoc(paragraphs=paragraphs, tables=tables)


def dispatch_parse(ext: str, data: bytes) -> ParsedDoc:
    """按扩展名分发解析器；旧格式给出明确另存提示。"""
    if ext == ".docx":
        return parse_docx(data)
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(data)
    if ext == ".pdf":
        return parse_pdf(data)
    if ext == ".doc":
        raise UnsupportedFormatError(
            "暂不支持旧版 .doc 格式，请将文件另存为 .docx 后重新上传"
        )
    if ext == ".xls":
        raise UnsupportedFormatError(
            "暂不支持旧版 .xls 格式，请将文件另存为 .xlsx 后重新上传"
        )
    if ext == ".pptx":
        raise UnsupportedFormatError(
            "暂不支持 .pptx 格式，请将文件导出为 .pdf 后重新上传"
        )
    raise UnsupportedFormatError(f"不支持的文件扩展名: {ext}")


# ---------- 文本切片 ----------


def chunk_text(
    paragraphs: list[str],
    target_size: int = CHUNK_TARGET_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[str]:
    """段落贪心聚合切片：单段 ≤ target；聚合超限时开新块并携带前块尾部 overlap。

    - 超长段落硬切（滑窗步长 target-overlap）
    - 相邻块 overlap：取前块末尾 ≤overlap 字符拼到新块头部（保证语义连续）
    """
    chunks: list[str] = []
    buffer = ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        # 超长段落硬切
        while len(para) > target_size:
            if buffer:
                chunks.append(buffer)
                buffer = ""
            chunks.append(para[:target_size])
            para = para[target_size - overlap :]
        if not para:
            continue
        if not buffer:
            buffer = para
        elif len(buffer) + 1 + len(para) <= target_size:
            buffer = f"{buffer}\n{para}"
        else:
            # 开新块：携带前块尾部（不超过剩余空间，保证新块 ≤ target）
            tail_limit = target_size - len(para) - 1
            tail = buffer[-overlap:] if tail_limit >= overlap else ""
            if tail:
                tail = tail.lstrip()
            chunks.append(buffer)
            buffer = f"{tail}\n{para}" if tail else para
    if buffer:
        chunks.append(buffer)
    return [c for c in (chunk.strip() for chunk in chunks) if c]


# ---------- 列映射与结构化抽取 ----------


def _normalize_header(header: str) -> str:
    return header.strip().lower().replace(" ", "").replace("\u3000", "")


def match_column(header: str, column_map: dict[str, str]) -> str | None:
    """表头 → 字段名匹配（归一化后查映射表）。"""
    return column_map.get(_normalize_header(header))


def _split_list_value(value: Any) -> list:
    """宽松解析列表字段：list 直返 > JSON 数组 > 分隔符切分（, ; ，；、换行）。"""
    if isinstance(value, list):
        return value
    text = str(value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return parsed
    except (ValueError, TypeError):
        pass
    return [p.strip() for p in re.split(r"[,;，；、\n]+", text) if p.strip()]


def _parse_dict_value(value: Any) -> dict:
    """宽松解析 dict 字段（环境 variables）：dict 直返 > JSON 对象 > k=v / k：v 键值对。"""
    if isinstance(value, dict):
        return value
    text = str(value or "").strip()
    if not text:
        return {}
    try:
        parsed = json.loads(text)
        if isinstance(parsed, dict):
            return parsed
    except (ValueError, TypeError):
        pass
    result: dict[str, str] = {}
    for part in re.split(r"[,;，；\n]+", text):
        pair = re.split(r"[=：:]", part, maxsplit=1)
        if len(pair) == 2 and pair[0].strip():
            result[pair[0].strip()] = pair[1].strip()
    return result


def _parse_number(value: str) -> float | None:
    """宽松解析数值：容忍百分号/千分位/单位后缀（如 500、95%、1,200、120ms）。"""
    text = (value or "").strip().replace(",", "")
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if match is None:
        return None
    try:
        return float(match.group())
    except ValueError:
        return None


def _trunc(value: str, limit: int) -> str:
    return (value or "").strip()[:limit]


def extract_inventory_rows(
    tables: list[list[list[str]]], column_map: dict[str, str]
) -> InventoryExtraction:
    """从表格提取行字典列表。

    - 表头行 = 每张表中首个命中 ≥2 个映射列的行；未识别到表头的表跳过并告警
    - 数据行仅保留映射到字段的非空值（row 字典值统一 str）
    - 出现但未匹配的表头记入 unmatched_columns（供 D5 列映射修正接口展示）
    """
    extraction = InventoryExtraction()
    seen_unmatched: set[str] = set()
    for table_index, table in enumerate(tables, start=1):
        header_index = -1
        field_by_col: dict[int, str] = {}
        for row_index, row in enumerate(table):
            col_map: dict[int, str] = {}
            unmatched_in_row: list[str] = []
            for col_index, cell in enumerate(row):
                if not cell:
                    continue
                matched = match_column(cell, column_map)
                if matched:
                    col_map[col_index] = matched
                else:
                    unmatched_in_row.append(cell)
            if len(col_map) >= 2:
                header_index = row_index
                field_by_col = col_map
                for header in unmatched_in_row:
                    if header not in seen_unmatched:
                        seen_unmatched.add(header)
                        extraction.unmatched_columns.append(header)
                break
            # 表头前的散行不产生 unmatched 噪音，直接忽略
        if header_index < 0:
            extraction.warnings.append(f"第 {table_index} 张表未识别到表头行，已跳过")
            continue
        for row in table[header_index + 1 :]:
            values: dict[str, str] = {}
            for col_index, cell in enumerate(row):
                if col_index in field_by_col and cell:
                    values[field_by_col[col_index]] = cell
            if values:
                extraction.rows.append(values)
    return extraction


def _build_environment_row(
    project_id: int, values: dict[str, str], warnings: list[str], row_label: str
) -> Environment | None:
    """行字典 → Environment ORM（env_code 缺失跳过并告警）。"""
    env_code = _trunc(values.get("env_code", ""), 64)
    if not env_code:
        warnings.append(f"{row_label}缺少环境编码（env_code），已跳过")
        return None
    name = _trunc(values.get("name", ""), 128) or env_code
    return Environment(
        project_id=project_id,
        name=name,
        env_code=env_code,
        base_url=_trunc(values.get("base_url", ""), 512),
        hosts=_split_list_value(values.get("hosts", "")),
        db_connections=_split_list_value(values.get("db_connections", "")),
        middleware_info=_split_list_value(values.get("middleware_info", "")),
        variables=_parse_dict_value(values.get("variables", "")),
        description=_trunc(values.get("description", ""), 512),
    )


def _build_transaction_row(
    project_id: int,
    values: dict[str, str],
    warnings: list[str],
    row_label: str,
    script_id_by_name: dict[str, int],
) -> Transaction | None:
    """行字典 → Transaction ORM（txn_code 缺失跳过；默认脚本按名称关联）。"""
    txn_code = _trunc(values.get("txn_code", ""), 64)
    if not txn_code:
        warnings.append(f"{row_label}缺少交易码（txn_code），已跳过")
        return None
    default_script_id: int | None = None
    script_name = values.get("default_script", "").strip()
    if script_name:
        default_script_id = script_id_by_name.get(script_name)
        if default_script_id is None:
            warnings.append(
                f"{row_label}默认脚本「{script_name}」在项目内不存在，未关联"
            )
    sla_tps = _parse_number(values.get("sla_tps", ""))
    sla_p95_ms = _parse_number(values.get("sla_p95_ms", ""))
    sla_error_rate = _parse_number(values.get("sla_error_rate", ""))
    return Transaction(
        project_id=project_id,
        name=_trunc(values.get("name", ""), 128) or txn_code,
        txn_code=txn_code,
        default_script_id=default_script_id,
        sla_tps=sla_tps,
        sla_p95_ms=int(sla_p95_ms) if sla_p95_ms is not None else None,
        sla_error_rate=sla_error_rate,
        description=_trunc(values.get("description", ""), 512),
    )


def _dedupe_rows(
    rows: list[Environment] | list[Transaction],
    key_attr: str,
    existing_keys: set[str],
    warnings: list[str],
    key_label: str,
) -> list[Any]:
    """项目内 + 文件内按键去重：重复跳过并告警（避免撞唯一约束）。"""
    kept = []
    seen_in_file: set[str] = set()
    for row in rows:
        key = getattr(row, key_attr)
        if key in existing_keys or key in seen_in_file:
            warnings.append(f"{key_label}={key} 已存在，该行跳过")
            continue
        seen_in_file.add(key)
        kept.append(row)
    return kept


def _table_text_paragraphs(tables: list[list[list[str]]]) -> list[str]:
    """表格 → 伪段落文本（供非清单类文档的表格内容进向量库）。"""
    paragraphs = []
    for table in tables:
        lines = [" | ".join(cells) for cells in table if cells]
        if lines:
            paragraphs.append("\n".join(lines))
    return paragraphs


# ---------- 解析入口与状态机 ----------


async def run_asset_parse(asset_id: int, session_factory=None) -> None:
    """解析入口（APScheduler job / retry 端点共用）。

    自带独立会话（调度器上下文无请求级会话）；任何异常兜底置 FAILED，
    不向调度器外抛（调度器 job 异常仅进日志）。
    session_factory 参数供测试注入内存库会话工厂。
    """
    factory = session_factory or SessionLocal
    try:
        async with factory() as db:
            await _parse_asset(db, asset_id)
    except Exception:  # noqa: BLE001
        logger.exception(f"资产 {asset_id} 解析管道异常")
        # _parse_asset 内部已兜底 FAILED；此处防其自身 DB 异常导致状态卡 PARSING
        try:
            async with factory() as db:
                await db.execute(
                    update(Asset)
                    .where(Asset.id == asset_id)
                    .values(
                        status=AssetStatus.FAILED.value,
                        parse_meta={"error": "解析管道异常（详见服务端日志）"},
                    )
                )
                await db.commit()
        except Exception:  # noqa: BLE001
            logger.exception(f"资产 {asset_id} 兜底置 FAILED 失败")


async def _parse_asset(db: AsyncSession, asset_id: int) -> None:
    """单资产解析主流程：CAS 抢占 PARSING → 解析 → READY/FAILED。

    CAS（status != parsing 才置 parsing）保证重复投递幂等：
    并发/重复 job 第二次进入时 rowcount=0 直接返回。
    """
    claimed = await db.execute(
        update(Asset)
        .where(Asset.id == asset_id, Asset.status != AssetStatus.PARSING.value)
        .values(status=AssetStatus.PARSING.value, parse_meta={})
    )
    await db.commit()
    if claimed.rowcount == 0:
        logger.warning(f"资产 {asset_id} 不存在或已在解析中，跳过本次解析")
        return

    asset = (
        await db.execute(select(Asset).where(Asset.id == asset_id))
    ).scalar_one_or_none()
    if asset is None:
        return

    meta: dict[str, Any] = {
        "chunks": 0,
        "indexed": False,
        "environments": 0,
        "transactions": 0,
        "warnings": [],
        "unmatched_columns": [],
    }
    try:
        if not asset.file_key:
            raise RuntimeError("资产文件对象缺失（file_key 为空），无法解析")
        data = await storage.get_object_bytes(asset.file_key)
        ext = Path(asset.filename or "").suffix.lower()
        parsed = await asyncio.to_thread(dispatch_parse, ext, data)

        # 路径 (b)：结构化抽取（仅清单类资产）
        column_map = _INVENTORY_COLUMN_MAPS.get(asset.asset_type)
        if column_map is not None:
            await _extract_structured(db, asset, parsed, column_map, meta)

        # 路径 (a)：文本切片 → embedding → 向量库（清单类资产表格已结构化，不再入向量）
        paragraphs = list(parsed.paragraphs)
        if column_map is None:
            paragraphs.extend(_table_text_paragraphs(parsed.tables))
        if paragraphs:
            chunks = chunk_text(paragraphs)
            meta["chunks"] = len(chunks)
            await _index_chunks(asset, chunks, meta)
        elif column_map is None:
            meta["warnings"].append("文档未提取到有效文本内容")

        await db.execute(
            update(Asset)
            .where(Asset.id == asset_id)
            .values(status=AssetStatus.READY.value, parse_meta=meta)
        )
        await db.commit()
        logger.info(
            f"资产 {asset_id} 解析完成: status=ready chunks={meta['chunks']} "
            f"indexed={meta['indexed']} environments={meta['environments']} "
            f"transactions={meta['transactions']}"
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception(f"资产 {asset_id} 解析失败")
        await db.rollback()
        meta["error"] = str(exc)[:500]
        await db.execute(
            update(Asset)
            .where(Asset.id == asset_id)
            .values(status=AssetStatus.FAILED.value, parse_meta=meta)
        )
        await db.commit()


async def _extract_structured(
    db: AsyncSession,
    asset: Asset,
    parsed: ParsedDoc,
    column_map: dict[str, str],
    meta: dict[str, Any],
) -> None:
    """表格行 → 环境/交易表：应用层去重后落库（唯一键冲突行跳过并告警）。"""
    extraction = extract_inventory_rows(parsed.tables, column_map)
    meta["unmatched_columns"] = extraction.unmatched_columns
    warnings = extraction.warnings

    if asset.asset_type == AssetType.ENV_INVENTORY.value:
        candidates = [
            row
            for row in (
                _build_environment_row(
                    asset.project_id, values, warnings, f"第{i + 1}行"
                )
                for i, values in enumerate(extraction.rows)
            )
            if row is not None
        ]
        existing_codes = set(
            (
                await db.execute(
                    select(Environment.env_code).where(
                        Environment.project_id == asset.project_id
                    )
                )
            ).scalars()
        )
        kept = _dedupe_rows(
            candidates, "env_code", existing_codes, warnings, "环境编码"
        )
        db.add_all(kept)
        meta["environments"] = len(kept)
    else:
        script_rows = (
            await db.execute(
                select(Script.id, Script.name).where(
                    Script.project_id == asset.project_id
                )
            )
        ).all()
        script_id_by_name = {name: sid for sid, name in script_rows}
        candidates = [
            row
            for row in (
                _build_transaction_row(
                    asset.project_id,
                    values,
                    warnings,
                    f"第{i + 1}行",
                    script_id_by_name,
                )
                for i, values in enumerate(extraction.rows)
            )
            if row is not None
        ]
        existing_codes = set(
            (
                await db.execute(
                    select(Transaction.txn_code).where(
                        Transaction.project_id == asset.project_id
                    )
                )
            ).scalars()
        )
        kept = _dedupe_rows(candidates, "txn_code", existing_codes, warnings, "交易码")
        db.add_all(kept)
        meta["transactions"] = len(kept)

    meta["warnings"].extend(warnings)
    # 截断告警列表，防止超大清单把 parse_meta 撑爆
    del meta["warnings"][50:]


async def _index_chunks(asset: Asset, chunks: list[str], meta: dict[str, Any]) -> None:
    """切片向量化并写入向量库。

    - 未配置 Embedding：降级跳过（indexed=false），结构化抽取不受影响
    - 已配置但调用失败：抛出 → 上层标记 FAILED（含重试耗尽信息）
    """
    settings = get_settings()
    if not (settings.embedding_api_key and settings.embedding_model):
        meta["warnings"].append("Embedding 未配置，文本切片未入向量库")
        return
    client = EmbeddingClient()
    vectors = await client.embed(chunks)
    store = get_vector_store()
    docs = [
        ChunkDoc(
            point_id=point_id_for(asset.id, index),
            asset_id=asset.id,
            project_id=asset.project_id,
            chunk_index=index,
            text_chunk=text,
            embedding=vector,
            asset_type=asset.asset_type,
            source_type="asset",
            source_ref=asset.file_key,
        )
        for index, (text, vector) in enumerate(zip(chunks, vectors, strict=True))
    ]
    await store.upsert_chunks(docs)
    meta["indexed"] = True


def schedule_asset_parse(asset_id: int, delay_seconds: float = 2.0) -> None:
    """上传/重试后延迟投递解析任务（给上传响应留出返回时间）。"""
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    from app.services.scheduler import enqueue_date_job

    settings = get_settings()
    run_date = datetime.now(ZoneInfo(settings.timezone)) + timedelta(
        seconds=delay_seconds
    )
    enqueue_date_job(
        run_asset_parse,
        job_key=f"asset-parse-{asset_id}",
        run_date=run_date,
        args=[asset_id],
    )
